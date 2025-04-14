# shape_analysis.py
import cv2
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import PatternFill
import os
import sys
from skimage import filters
from skimage.measure import label, regionprops
from skimage.morphology import closing, footprint_rectangle
import json  # ✅ Add this at the top if not already imported
MARGIN = 0
PADDING = 10
LOCAL_PAD = 30

def analyze_shapes_and_export(min_area=10, output_dir="static"):
    # --- Load images ---
    before_img = cv2.imread(os.path.join(output_dir, "FranjeBefore.jpg"))
    after_img = cv2.imread(os.path.join(output_dir, "Franjeafter.jpg"))
    assert before_img is not None and after_img is not None, "Frangi image load failed!"

    original_before = cv2.imread(os.path.join(output_dir, "OriginalBefor.jpg"))
    original_after = cv2.imread(os.path.join(output_dir, "OriginalAfter.jpg"))
    assert original_before is not None and original_after is not None, "Original image load failed!"

    gray_before = cv2.cvtColor(before_img, cv2.COLOR_BGR2GRAY)
    gray_after = cv2.cvtColor(after_img, cv2.COLOR_BGR2GRAY)

    # --- Segment clusters in BEFORE ---
    thresh_before = filters.threshold_otsu(gray_before)
    binary_before = gray_before > thresh_before
    binary_before_closed = closing(binary_before, footprint_rectangle((5, 5)))
    label_before = label(binary_before_closed)
    regions = sorted(regionprops(label_before), key=lambda r: r.bbox[0])

    os.makedirs(os.path.join(output_dir, "cluster_crops"), exist_ok=True)
    data = []

    for idx, region in enumerate(regions):
        cluster_id = idx + 1
        minr, minc, maxr, maxc = expand_to_shape_bounds(region.bbox, binary_before_closed, gray_before.shape)

        crop_before = gray_before[minr:maxr, minc:maxc]
        crop_after = gray_after[minr:maxr, minc:maxc]

        count_before, vis_before = count_and_draw_shapes(crop_before, (0, 0, 255), min_area=min_area)
        count_after, vis_after = count_and_draw_shapes(crop_after, (255, 0, 0), min_area=min_area)

        difference = count_after - count_before
        percent_increase = 0.0
        if count_before > 0:
            percent_increase = round((difference / count_before) * 100, 2)
        elif count_after > 0:
            percent_increase = 100.0

        if difference > 0:
            change_type = "Gain"
        elif difference < 0:
            change_type = "Loss"
        else:
            change_type = "No Change"

        # --- Visuals ---
        zb_top = max(0, minr - LOCAL_PAD)
        zb_bottom = min(original_before.shape[0], maxr + LOCAL_PAD)
        zb_left = max(0, minc - LOCAL_PAD)
        zb_right = min(original_before.shape[1], maxc + LOCAL_PAD)

        zoomed_before = original_before[zb_top:zb_bottom, zb_left:zb_right].copy()
        zoomed_after = original_after[zb_top:zb_bottom, zb_left:zb_right].copy()

        cv2.rectangle(zoomed_before, (LOCAL_PAD, LOCAL_PAD), 
                      (zoomed_before.shape[1] - LOCAL_PAD, zoomed_before.shape[0] - LOCAL_PAD),
                      (0, 255, 255), 2)
        cv2.rectangle(zoomed_after, (LOCAL_PAD, LOCAL_PAD),
                      (zoomed_after.shape[1] - LOCAL_PAD, zoomed_after.shape[0] - LOCAL_PAD),
                      (0, 255, 255), 2)

        labeled_context = np.hstack((zoomed_before, zoomed_after))
        labeled_detection = np.hstack((vis_before, vis_after))

        h1, w1 = labeled_detection.shape[:2]
        h2, w2 = labeled_context.shape[:2]
        if w1 != w2:
            labeled_context = cv2.resize(labeled_context, (w1, h2), interpolation=cv2.INTER_AREA)

        combined = np.vstack((labeled_detection, labeled_context))
        cv2.imwrite(os.path.join(output_dir, "cluster_crops", f"Cluster_{cluster_id:02d}_shapes.png"), combined)

        data.append((cluster_id, count_before, count_after, difference, percent_increase, change_type))

    # --- DataFrame ---
    df = pd.DataFrame(data, columns=[
        "Cluster Index", "Before Shape Count", "After Shape Count", "Difference", "% Increase", "Change Type"
    ])
    excel_path = os.path.join(output_dir, "cluster_shape_summary.xlsx")
    df.to_excel(excel_path, index=False)

    # --- Excel Formatting ---
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    yellow = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        diff = row[3].value
        change = row[5].value
        if diff > 0:
            row[3].fill = green
        elif diff < 0:
            row[3].fill = red

        if change == "Gain":
            row[5].fill = green
        elif change == "Loss":
            row[5].fill = red
        elif change == "No Change":
            row[5].fill = yellow

    wb.save(excel_path)
    wb.close()

    # --- Save status ---
    with open(os.path.join(output_dir, "analysis_status.txt"), "w") as f:
        f.write("done")
        # --- Save total stats for web display ---
   
    # --- Print summary ---
    total_before = df["Before Shape Count"].sum()
    total_after = df["After Shape Count"].sum()
    total_diff = total_after - total_before
    percent_total_change = round((total_diff / total_before) * 100, 2) if total_before else 100.0

    # ✅ Gain-only percent increase
    gain_rows = df[df["Difference"] > 0]
    gain_before = gain_rows["Before Shape Count"].sum()
    gain_after = gain_rows["After Shape Count"].sum()
    gain_diff = gain_after - gain_before
    percent_gain_only = round((gain_diff / gain_before) * 100, 2) if gain_before else 0.0

    # ✅ Percentages of cluster change types
    total_clusters = len(df)
    gain_percent = round((df["Change Type"] == "Gain").sum() / total_clusters * 100, 2)
    loss_percent = round((df["Change Type"] == "Loss").sum() / total_clusters * 100, 2)
    nochange_percent = round((df["Change Type"] == "No Change").sum() / total_clusters * 100, 2)

    print("\n ANALYSIS SUMMARY:")
    print(f"Total Before:  {total_before}")
    print(f"Total After:   {total_after}")
    print(f"Net Change:    {total_diff}")
    print(f"Percent Change (net): {percent_total_change:.2f}%")
    print(f"Percent Change (gains only): {percent_gain_only:.2f}%")
    print(f"% Clusters with Gain: {gain_percent}%, Loss: {loss_percent}%, No Change: {nochange_percent}%")

    # --- Save total stats for web display ---
    summary_path = os.path.join(output_dir, "summary.json")
    with open(summary_path, "w") as f:
        json.dump({
            "total_before": int(total_before),
            "total_after": int(total_after),
            "net_change": int(total_diff),
            "percent_change": percent_total_change,
            "gain_only_percent": percent_gain_only,
            "gain_percent": gain_percent,
            "loss_percent": loss_percent,
            "nochange_percent": nochange_percent
        }, f)

    # --- Also save a human-readable TXT summary ---
    results_txt_path = os.path.join(output_dir, "results.txt")
    with open(results_txt_path, "w") as f:
        f.write("ANALYSIS SUMMARY\n")

        # Load metadata from info.json if available
        try:
            with open(os.path.join(output_dir, "info.json")) as meta_file:
                info = json.load(meta_file)
            f.write(f"Before Image: {info.get('before_image', 'unknown')}\n")
            f.write(f"After Image: {info.get('after_image', 'unknown')}\n")
            f.write(f"Cropped: {'Yes' if info.get('cropped') else 'No'}\n")
            if info.get("cropped") and "crop_box" in info:
                box = info["crop_box"]
                f.write(f"Crop Area: x=({box.get('x_start')} - {box.get('x_end')}), y=({box.get('y_start')} - {box.get('y_end')})\n")
            f.write(f"Min Shape Area: {info.get('min_area', 'unknown')}\n\n")
        except Exception as e:
            f.write("Info.json not found or unreadable\n\n")

        f.write(f"Total Before: {total_before}\n")
        f.write(f"Total After: {total_after}\n")
        f.write(f"Net Change: {total_diff}\n")
        f.write(f"Percent Change (net): {percent_total_change:.2f}%\n")
        f.write(f"Percent Change (gains only): {percent_gain_only:.2f}%\n")
        f.write(f"% Clusters with Gain: {gain_percent}%\n")
        f.write(f"% Clusters with No Change: {nochange_percent}%\n")
        f.write(f"% Clusters with Loss: {loss_percent}%\n")



def expand_to_shape_bounds(bbox, binary_mask, max_shape):
    minr, minc, maxr, maxc = bbox
    region_mask = np.zeros_like(binary_mask, dtype=np.uint8)
    region_mask[minr:maxr, minc:maxc] = 1
    shape_mask = binary_mask * region_mask
    labeled = label(shape_mask)
    props = regionprops(labeled)
    if not props:
        return minr, minc, maxr, maxc
    all_coords = np.concatenate([p.coords for p in props])
    r_vals = all_coords[:, 0]
    c_vals = all_coords[:, 1]
    minr_adj = max(0, np.min(r_vals) - MARGIN - PADDING)
    minc_adj = max(0, np.min(c_vals) - MARGIN - PADDING)
    maxr_adj = min(max_shape[0], np.max(r_vals) + MARGIN + PADDING)
    maxc_adj = min(max_shape[1], np.max(c_vals) + MARGIN + PADDING)
    return minr_adj, minc_adj, maxr_adj, maxc_adj

def count_and_draw_shapes(image_crop, color, min_area=20):
    vis = cv2.cvtColor(image_crop, cv2.COLOR_GRAY2BGR)
    blurred = cv2.GaussianBlur(image_crop, (3, 3), 0)
    thresh_val = filters.threshold_otsu(blurred)
    binary = (blurred > thresh_val).astype(np.uint8) * 255
    contours, _ = cv2.findContours(binary, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    valid_contours = [cnt for cnt in contours if cv2.contourArea(cnt) >= min_area]
    for cnt in valid_contours:
        x, y, w, h = cv2.boundingRect(cnt)
        cv2.rectangle(vis, (x, y), (x + w, y + h), color, 1)
    return len(valid_contours), vis

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--min_area', type=int, default=10)
    parser.add_argument('--output_dir', type=str, default="static")
    args = parser.parse_args()
    analyze_shapes_and_export(min_area=args.min_area, output_dir=args.output_dir)
